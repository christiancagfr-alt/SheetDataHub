from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from sheet_hub.config_store import ConfigStore
from sheet_hub.database import AggregateDatabase, SourceCache
from sheet_hub.engine import (
    DATE_HEADER_HINTS,
    DataEngine,
    cell_has_d,
    compare_periods,
    is_chart_header,
    is_session_header,
    is_stat_numeric_header,
    last_n_days_ranges,
    list_chart_headers,
    list_matching_headers,
    list_name_headers,
    month_compare_ranges,
    parse_date,
    parse_number,
    pick_header,
    week_compare_ranges,
)
from sheet_hub.models import Record, SourceConfig
from sheet_hub.source_reader import (
    CredentialPool,
    SourceReader,
    _google_download_host_ok,
    canonicalize,
    choose_sheets,
    google_retry,
    parse_schema_lines,
)
from sheet_hub.ui import DEFAULT_QUERY_RESULT_FIELDS, is_phone_data_table, query_result_headers
from sheet_hub.version import (
    APP_VERSION,
    download_release_installer,
    fetch_latest_release,
    installer_url_allowed,
    is_newer,
    parse_version,
)


class RuleTests(unittest.TestCase):
    def test_google_download_host_allowlist(self):
        self.assertTrue(_google_download_host_ok("https://docs.google.com/spreadsheets/d/abc/export?format=xlsx"))
        self.assertTrue(_google_download_host_ok("https://doc-00-00-docs.googleusercontent.com/file"))
        self.assertFalse(_google_download_host_ok("https://evil.example/file.xlsx"))
        self.assertFalse(_google_download_host_ok("https://docs.google.com.evil.test/file"))

    def test_empty_include_runs_all_except_excluded(self):
        selected, missing = choose_sheets(["订单", "统计", "客户"], [], ["统计"])
        self.assertEqual(selected, ["订单", "客户"])
        self.assertEqual(missing, [])

    def test_include_only_and_exclude_wins(self):
        selected, missing = choose_sheets(["订单A", "订单B"], ["订单A", "订单B", "不存在"], ["订单B"])
        self.assertEqual(selected, ["订单A"])
        self.assertEqual(missing, ["不存在"])

    def test_version_compare(self):
        self.assertEqual(parse_version("v1.2.0"), (1, 2, 0))
        self.assertTrue(is_newer("1.2.1", "1.2.0"))
        self.assertFalse(is_newer("1.2.0", "1.2.0"))
        self.assertFalse(is_newer("1.1.9", APP_VERSION))
        self.assertEqual(parse_schema_lines("专页ID\n姓名\n\n号码\n"), ["专页ID", "姓名", "", "号码"])

    def test_space_separated_sheet_exclusions_are_supported(self):
        selected, _ = choose_sheets(["数据", "Index", "清理", "备份链接"], [], ["Index 清理 备份链接"])
        self.assertEqual(selected, ["数据"])

    def test_alias_mapping(self):
        result = canonicalize({"手机号": 13800138000, "渠道": "广告"}, {"号码": ["手机号"], "来源": ["渠道"]})
        self.assertEqual(result["号码"], "13800138000")
        self.assertEqual(result["来源"], "广告")

    def test_date_parsing(self):
        self.assertEqual(parse_date("2026/09/13"), date(2026, 9, 13))
        self.assertEqual(parse_date("2026年9月13日"), date(2026, 9, 13))
        self.assertIsNone(parse_date("not-a-date"))

    def test_update_release_selects_and_downloads_installer(self):
        class FakeJsonResponse:
            status_code = 200

            def raise_for_status(self):
                return None

            def json(self):
                return {
                    "tag_name": "v9.9.9",
                    "html_url": "https://example.test/release",
                    "assets": [{
                        "name": "SheetDataHub-Setup-v9.9.9.exe",
                        "browser_download_url": "https://github.com/christiancagfr-alt/SheetDataHub/releases/download/v9.9.9/SheetDataHub-Setup-v9.9.9.exe",
                        "size": 1024 * 1024 + 2,
                    }],
                }

        with patch("sheet_hub.version.requests.get", return_value=FakeJsonResponse()):
            info = fetch_latest_release()
        self.assertTrue(installer_url_allowed(info["installer_url"]))
        self.assertIn("github.com/", info["installer_url"])
        self.assertFalse(installer_url_allowed("https://example.test/setup.exe"))

        payload = b"MZ" + (b"x" * (1024 * 1024))

        class FakeDownloadResponse:
            url = info["installer_url"]

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def raise_for_status(self):
                return None

            def iter_content(self, chunk_size):
                yield payload

        with tempfile.TemporaryDirectory() as directory, patch(
            "sheet_hub.version.requests.get", return_value=FakeDownloadResponse()
        ):
            path = download_release_installer(info, directory)
            self.assertEqual(path.read_bytes()[:2], b"MZ")
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(RuntimeError):
                download_release_installer(
                    {"installer_url": "https://example.test/setup.exe", "version": "1.0.0"},
                    directory,
                )
        class FakeRedirectResponse(FakeDownloadResponse):
            url = "https://example.test/evil.exe"

        with tempfile.TemporaryDirectory() as directory, patch(
            "sheet_hub.version.requests.get", return_value=FakeRedirectResponse()
        ):
            with self.assertRaises(RuntimeError):
                download_release_installer(info, directory)

    def test_credential_pool_rotates_on_429(self):
        class QuotaError(Exception):
            def __init__(self):
                self.response = SimpleNamespace(status_code=429)

        used: list[str] = []

        def factory(path):
            return path

        def operation(client):
            used.append(Path(client).name)
            if len(used) == 1:
                raise QuotaError()
            return "ok"

        with tempfile.TemporaryDirectory() as directory, patch("sheet_hub.source_reader.time.sleep"):
            first = Path(directory) / "a.json"
            second = Path(directory) / "b.json"
            first.write_text("{}", encoding="utf-8")
            second.write_text("{}", encoding="utf-8")
            pool = CredentialPool([str(first), str(second)], client_factory=factory)
            self.assertEqual(pool.call(operation), "ok")
        self.assertEqual(used, ["a.json", "b.json"])

    def test_credential_paths_migrate_legacy_setting(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", r"C:\keys\one.json")
            self.assertEqual(store.list_credential_paths(), [r"C:\keys\one.json"])
            store.set_credential_paths([r"C:\keys\one.json", r"C:\keys\two.json"])
            self.assertEqual(store.list_credential_paths(), [r"C:\keys\one.json", r"C:\keys\two.json"])
            self.assertEqual(store.get("credential_path"), r"C:\keys\one.json")

    def test_google_429_is_retried(self):
        class QuotaError(Exception):
            response = SimpleNamespace(status_code=429)

        calls = {"count": 0}

        def operation():
            calls["count"] += 1
            if calls["count"] == 1:
                raise QuotaError()
            return "ok"

        with patch("sheet_hub.source_reader.time.sleep"), patch("sheet_hub.source_reader.random.random", return_value=0):
            self.assertEqual(google_retry(operation), "ok")
        self.assertEqual(calls["count"], 2)

    def test_phone_matching_normalizes_common_formats(self):
        self.assertEqual(DataEngine._match_value("号码", "258-851-758692"), "258851758692")
        self.assertEqual(DataEngine._match_value("号码", "258851758692.0"), "258851758692")
        self.assertEqual(DataEngine._match_value("号码", "2.58851758692E+11"), "258851758692")

    def test_existing_aliases_are_migrated(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(directory)
            aliases = store.get("field_aliases")
            aliases["号码"] = ["号码", "手机号"]
            store.set("field_aliases", aliases)
            reopened = ConfigStore(directory)
            self.assertIn("手机号码", reopened.get("field_aliases")["号码"])

    def test_query_source_is_remembered(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(directory)
            self.assertEqual(store.get("query_source"), "extract")
            store.set("query_source", "direct")
            store.set("google_output_url", "https://docs.google.com/spreadsheets/d/abc")
            store.set("google_output_sheet", "提取结果")
            reopened = ConfigStore(directory)
            self.assertEqual(reopened.get("query_source"), "direct")
            self.assertEqual(reopened.get("google_output_sheet"), "提取结果")
            self.assertTrue(reopened.get("query_exact"))
            self.assertFalse(reopened.get("query_fuzzy"))
            self.assertFalse(reopened.get("query_date_enabled"))

    def test_query_result_headers_keep_phone_format(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(directory)
            source_headers = ["预交表汇总", "见证状态", "交教会日期", "摸底/推广", "线索电话号码"]
            self.assertEqual(
                query_result_headers(store, "direct", "线索电话号码", source_headers, "交教会"),
                source_headers,
            )
            self.assertEqual(
                query_result_headers(store, "direct", "摸底/推广", source_headers, "交教会"),
                source_headers,
            )
            self.assertEqual(
                query_result_headers(store, "direct", "号码", ["专页ID", "姓名", "号码"], "号码表"),
                DEFAULT_QUERY_RESULT_FIELDS,
            )
            self.assertFalse(is_phone_data_table(source_headers, "交教会"))
            self.assertTrue(is_phone_data_table(["专页ID", "号码"], "交教会"))

    def test_source_cache_and_selected_sync(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            path = root / "club.xlsx"
            book = openpyxl.Workbook()
            book.active.title = "数据"
            book.active.append(["加友途径4", "贴文ID"])
            book.active.append(["2081专页后台", "p1"])
            book.save(path)
            store.save_source(SourceConfig(
                "club", "交教会", str(path),
                column_schema_enabled=True,
                column_schema=[
                    {"name": "加友途径4", "column": "A", "enabled": True},
                    {"name": "贴文ID", "column": "B", "enabled": True},
                ],
            ))
            engine = DataEngine(store)
            synced = engine.sync(["club"], write_local_db=True)
            self.assertEqual(synced["sources"], 1)
            self.assertTrue(engine.cache.has("club"))
            self.assertEqual(engine.list_query_fields("direct", source_id="club"), ["加友途径4", "贴文ID"])
            found = engine.query("加友途径4", "2081专页后台", source="direct", source_id="club")
            self.assertEqual(found[0].values["贴文ID"], "p1")


class DatabaseTests(unittest.TestCase):
    def test_clear_extract_cache_only_removes_local_dedup_history(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.mark_extracted(["k1", "k2"], "目标")
            store.log("INFO", "测试", "保留日志")
            self.assertEqual(store.count_extracted(), 2)
            self.assertEqual(store.clear_extracted(), 2)
            self.assertEqual(store.count_extracted(), 0)
            self.assertFalse(store.was_extracted("k1"))
            self.assertEqual(len(store.read_logs()), 1)

    def test_export_import_config_moves_sources_and_settings_only(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_store = ConfigStore(root / "source")
            source_store.set("global_excludes", ["不要跑"])
            source_store.set("extract_dedup_fields", "号码,日期")
            source_store.set("column_schema_enabled", True)
            source_store.set("column_schema", [{"name": "手机号码", "column": "H", "enabled": True}])
            source_store.set("credential_path", r"C:\secrets\service-account.json")
            source_store.save_source(SourceConfig(
                "s1", "主表", "https://docs.google.com/spreadsheets/d/abc",
                include_sheets=["订单"], exclude_sheets=["统计"], header_row=1,
                credential_path=r"C:\secrets\service-account.json",
                column_schema_enabled=True,
                column_schema=[{"name": "日期", "column": "I", "enabled": True}],
            ))
            source_store.mark_extracted(["old"], "目标")
            payload = source_store.export_config()

            target_store = ConfigStore(root / "target")
            target_store.mark_extracted(["keep"], "目标")
            settings_count, source_count = target_store.import_config(payload)

            self.assertGreater(settings_count, 0)
            self.assertEqual(source_count, 1)
            self.assertEqual(target_store.get("global_excludes"), ["不要跑"])
            self.assertEqual(target_store.get("column_schema")[0]["name"], "手机号码")
            imported_source = target_store.load_sources()[0]
            self.assertEqual(imported_source.name, "主表")
            self.assertEqual(imported_source.include_sheets, ["订单"])
            self.assertFalse(target_store.was_extracted("old"))
            self.assertTrue(target_store.was_extracted("keep"))
            self.assertEqual(payload["settings"].get("credential_path"), "")
            self.assertEqual(payload["settings"].get("credential_paths"), [])
            self.assertEqual(payload["sources"][0].get("credential_path"), "")
            target_store.set_credential_paths([r"D:\local\sa.json"])
            target_store.import_config({
                **payload,
                "settings": {
                    **payload["settings"],
                    "credential_path": r"C:\stolen\sa.json",
                    "credential_paths": [r"C:\stolen\sa.json"],
                },
                "sources": [{**payload["sources"][0], "credential_path": r"C:\stolen\sa.json"}],
            })
            self.assertEqual(target_store.list_credential_paths(), [r"D:\local\sa.json"])
            self.assertEqual(target_store.load_sources()[0].credential_path, "")

    def test_sharding_and_query(self):
        with tempfile.TemporaryDirectory() as directory:
            db = AggregateDatabase(directory, 1000)
            records = [Record("s", "源", "g", "订单", index, {"号码": str(index)}, str(index)) for index in range(1001)]
            rows, databases = db.replace_all(records)
            self.assertEqual(rows, 1001)
            self.assertEqual(databases, 2)
            self.assertEqual(db.query("号码", "1000")[0].row_number, 1000)

    def test_extract_dedup(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            engine = DataEngine(store)
            records = [Record("s", "源", "g", "订单", 2, {"日期": "2026-09-13", "号码": "123"}, "hash")]
            engine.database.replace_all(records)
            output1 = root / "one.xlsx"
            output2 = root / "two.xlsx"
            first = engine.extract("日期", date(2026, 9, 1), date(2026, 9, 30), output1, ["号码", "日期"])
            second = engine.extract("日期", date(2026, 9, 1), date(2026, 9, 30), output2, ["号码", "日期"])
            self.assertEqual(first["written"], 1)
            self.assertEqual(second["written"], 0)
            self.assertTrue(output2.exists())

    def test_batch_query_keeps_missing_numbers(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1008-李薇", 2, {"号码": "123", "名字": "李薇", "日期": "2026-09-13"}, "h1")
            ])
            results = engine.query_many("号码", ["123", "999"])
            self.assertEqual(results[0][1].sheet_name, "1008-李薇")
            self.assertIsNone(results[1][1])

    def test_query_excludes_multiple_keywords(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "A", 2, {"号码": "123", "备注": "正常客户"}, "h1"),
                Record("s", "源", "g", "B", 3, {"号码": "123", "备注": "广告推广"}, "h2"),
                Record("s", "源", "g", "C", 4, {"号码": "123", "备注": "测试账号"}, "h3"),
            ])
            results = engine.query_many("号码", ["123"], exclude_keywords=["广告", "测试"])
            self.assertEqual([record.sheet_name for _, record in results if record], ["A"])
            none = engine.query_many("号码", ["123"], exclude_keywords=["正常", "广告", "测试"])
            self.assertIsNone(none[0][1])

    def test_parse_number_strips_commas_and_percent(self):
        self.assertEqual(parse_number("1,234.50"), 1234.5)
        self.assertEqual(parse_number("12%"), 12.0)
        self.assertIsNone(parse_number("广告"))
        self.assertIsNone(parse_number(""))

    def test_week_and_month_compare_ranges(self):
        reference = date(2026, 10, 7)
        (current_start, current_end), (previous_start, previous_end) = week_compare_ranges(reference)
        self.assertEqual((current_start, current_end), (date(2026, 10, 5), date(2026, 10, 7)))
        self.assertEqual((previous_start, previous_end), (date(2026, 9, 28), date(2026, 9, 30)))
        (current_start, current_end), (previous_start, previous_end) = month_compare_ranges(reference)
        self.assertEqual((current_start, current_end), (date(2026, 10, 1), date(2026, 10, 7)))
        self.assertEqual((previous_start, previous_end), (date(2026, 9, 1), date(2026, 9, 7)))
        (_cs, _ce), (previous_start, previous_end) = month_compare_ranges(date(2026, 3, 31))
        self.assertEqual((previous_start, previous_end), (date(2026, 2, 1), date(2026, 2, 28)))
        (current_start, current_end), (previous_start, previous_end) = last_n_days_ranges(reference, 7)
        self.assertEqual((current_start, current_end), (date(2026, 10, 1), date(2026, 10, 7)))
        self.assertEqual((previous_start, previous_end), (date(2026, 9, 24), date(2026, 9, 30)))
        (current_start, current_end), (previous_start, previous_end) = last_n_days_ranges(reference, 2)
        self.assertEqual((current_start, current_end), (date(2026, 10, 6), date(2026, 10, 7)))
        self.assertEqual((previous_start, previous_end), (date(2026, 10, 4), date(2026, 10, 5)))
        current, previous = compare_periods("days7", reference)
        self.assertEqual(current, (date(2026, 10, 1), date(2026, 10, 7)))
        current, previous = compare_periods("month", reference, compare=False)
        self.assertEqual(current, (date(2026, 10, 1), date(2026, 10, 7)))
        self.assertLess(previous[1], current[0])

    def test_custom_compare_periods_require_complete_ranges(self):
        with self.assertRaises(ValueError):
            compare_periods("custom", date(2026, 10, 7))
        current, previous = compare_periods(
            "custom",
            date(2026, 10, 7),
            date(2026, 10, 1),
            date(2026, 10, 7),
            date(2026, 9, 1),
            date(2026, 9, 7),
        )
        self.assertEqual(current, (date(2026, 10, 1), date(2026, 10, 7)))
        self.assertEqual(previous, (date(2026, 9, 1), date(2026, 9, 7)))

    def _analysis_records(self):
        def rec(row: int, name: str, day: str, score: str, note: str = "") -> Record:
            return Record("s", "源", "g", name, row, {"名字": name, "日期": day, "业绩": score, "备注": note}, f"h{row}")

        return [
            rec(2, "张三", "2026-10-05", "10"),
            rec(3, "张三", "2026-10-06", "20"),
            rec(4, "李四", "2026-10-05", "5"),
            rec(5, "李四", "2026-10-06", "5"),
            rec(6, "张三", "2026-09-28", "8"),
            rec(7, "广告号", "2026-10-06", "100", "广告推广"),
        ]

    def test_analyze_team_week_compare_exclude_and_daily_growth(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(self._analysis_records())
            result = engine.analyze(
                source="aggregate",
                date_field="日期",
                name_field="名字",
                metric_field="业绩",
                exclude_keywords=["广告"],
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
        self.assertEqual(result.scope, "全部队别 · 全部人员")
        self.assertEqual(result.current.count, 4)
        self.assertEqual(result.current.total, 40)
        self.assertEqual(result.current.average, 10)
        self.assertEqual(result.previous.count, 1)
        self.assertEqual(result.previous.total, 8)
        self.assertEqual(result.delta_total, 32)
        people = {item.name: item for item in result.people}
        self.assertEqual(people["张三"].total, 30)
        self.assertEqual(people["张三"].count, 2)
        self.assertEqual(people["张三"].previous_count, 1)
        self.assertEqual(people["李四"].total, 10)
        self.assertEqual(people["李四"].previous_count, 0)
        self.assertNotIn("广告号", people)
        daily = {point.day: point for point in result.current.daily}
        self.assertEqual(daily["2026-10-05"].total, 15)
        self.assertEqual(daily["2026-10-06"].total, 25)
        self.assertEqual(daily["2026-10-06"].growth, 10)
        previous_daily = {point.day: point.count for point in result.previous.daily}
        self.assertEqual(previous_daily["2026-09-28"], 1)
        self.assertIn("2026-09-29", previous_daily)
        headers = {item.name: item for item in result.header_stats}
        self.assertEqual(headers["业绩"].total, 40)
        self.assertNotIn("日期", headers)

    def test_analyze_person_filter_and_count_metric(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(self._analysis_records())
            result = engine.analyze(
                source="aggregate",
                date_field="日期",
                name_field="名字",
                metric_field="",
                names=["张三"],
                exclude_keywords=["广告"],
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
        self.assertEqual(result.scope, "全部队别 · 张三")
        self.assertEqual(result.metric_field, "记录数")
        self.assertEqual(result.current.count, 2)
        self.assertEqual(result.current.total, 2)
        self.assertEqual([item.name for item in result.people], ["张三"])

    def test_analyze_custom_period_and_missing_person(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(self._analysis_records())
            result = engine.analyze(
                source="aggregate",
                date_field="日期",
                name_field="名字",
                metric_field="业绩",
                compare_mode="custom",
                current_start=date(2026, 10, 5),
                current_end=date(2026, 10, 6),
                previous_start=date(2026, 9, 28),
                previous_end=date(2026, 9, 28),
            )
            self.assertEqual(result.current.total, 140)
            self.assertEqual(result.previous.total, 8)
            with self.assertRaises(ValueError):
                engine.analyze(
                    source="aggregate",
                    date_field="日期",
                    name_field="名字",
                    names=["不存在的人"],
                    compare_mode="week",
                    reference_date=date(2026, 10, 6),
                )

    def test_pick_header_follows_schema_and_skips_lead_name(self):
        headers = ["队别", "状态", "交教会日期", "组别", "摸底/推广", "见证日期", "线索名字"]
        self.assertEqual(pick_header(headers, ("日期", "时间")), "交教会日期")
        self.assertEqual(pick_header(headers, ("队别", "队伍")), "队别")
        self.assertEqual(pick_header(headers, ("摸底/推广", "姓名", "名字"), ("线索",)), "摸底/推广")
        self.assertEqual(
            list_matching_headers(headers, DATE_HEADER_HINTS),
            ["交教会日期", "见证日期"],
        )
        self.assertFalse(is_stat_numeric_header("线索电话号码"))
        self.assertFalse(is_stat_numeric_header("贴文链接"))
        self.assertTrue(is_stat_numeric_header("业绩"))
        self.assertTrue(is_chart_header("加友途径4"))
        self.assertFalse(is_chart_header("线索电话号码"))
        self.assertEqual(
            list_chart_headers(headers + ["线索电话号码", "加友途径4", "贴文链接"]),
            ["状态", "组别", "加友途径4", "队别", "摸底/推广", "线索名字"],
        )
        self.assertEqual(list_name_headers(headers), ["摸底/推广", "线索名字"])
        self.assertTrue(is_session_header("第一场"))
        self.assertTrue(is_session_header("第八场"))
        self.assertFalse(is_session_header("市场"))
        self.assertTrue(cell_has_d("10.6 D 97min/笔聊73min"))
        self.assertTrue(cell_has_d("①D 2min"))
        self.assertFalse(cell_has_d("Não"))
        self.assertFalse(cell_has_d("10.2 1min"))
        self.assertFalse(cell_has_d("AND"))

    def test_analyze_team_and_name_are_independent(self):
        records = [
            Record("s", "交教会", "g", "浇灌", 2, {"队别": "安桑1队", "交教会日期": "2026-10-05", "摸底/推广": "心路Maria/张小川"}, "a"),
            Record("s", "交教会", "g", "浇灌", 3, {"队别": "安桑1队", "交教会日期": "2026-10-06", "摸底/推广": "安雨Anna/张小川3"}, "b"),
            Record("s", "交教会", "g", "浇灌", 4, {"队别": "华人队", "交教会日期": "2026-10-06", "摸底/推广": "心路Maria/张小川"}, "c"),
            Record("s", "交教会", "g", "浇灌", 5, {"队别": "安桑1队", "交教会日期": "2026-10-06", "摸底/推广": "智岩/阿黎"}, "d"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(records)
            whole_team = engine.analyze(
                source="aggregate",
                team="安桑1队",
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
            person = engine.analyze(
                source="aggregate",
                team="安桑1队",
                names=["张小川"],
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
            all_teams_person = engine.analyze(
                source="aggregate",
                names=["张小川"],
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
        self.assertEqual(whole_team.date_field, "交教会日期")
        self.assertEqual(whole_team.team_field, "队别")
        self.assertEqual(whole_team.name_field, "摸底/推广")
        self.assertEqual(whole_team.current.count, 3)
        self.assertEqual(whole_team.scope, "队别 安桑1队 · 整个队别")
        self.assertEqual(person.current.count, 2)
        self.assertEqual(person.scope, "队别 安桑1队 · 张小川")
        self.assertEqual(all_teams_person.current.count, 3)
        self.assertEqual(all_teams_person.scope, "全部队别 · 张小川")

    def test_analyze_uses_selected_date_column(self):
        records = [
            Record("s", "交教会", "g", "浇灌", 2, {
                "队别": "安桑1队",
                "交教会日期": "2026-10-05",
                "见证日期": "2026-09-28",
                "摸底/推广": "张小川",
                "线索电话号码": "15512345678",
            }, "a"),
            Record("s", "交教会", "g", "浇灌", 3, {
                "队别": "安桑1队",
                "交教会日期": "2026-10-06",
                "见证日期": "2026-10-06",
                "摸底/推广": "李四",
                "线索电话号码": "15587654321",
            }, "b"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(records)
            by_meet = engine.analyze(
                source="aggregate",
                date_field="交教会日期",
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
            by_witness = engine.analyze(
                source="aggregate",
                date_field="见证日期",
                compare_mode="week",
                reference_date=date(2026, 10, 6),
            )
        self.assertEqual(by_meet.date_field, "交教会日期")
        self.assertEqual(by_meet.current.count, 2)
        self.assertEqual(by_meet.previous.count, 0)
        self.assertEqual(by_witness.date_field, "见证日期")
        self.assertEqual(by_witness.current.count, 1)
        self.assertEqual(by_witness.previous.count, 1)
        self.assertEqual([point.day for point in by_meet.previous.daily], ["2026-09-28", "2026-09-29"])
        phones = {item.name: item for item in by_meet.header_stats}["线索电话号码"]
        self.assertFalse(phones.numeric)
        self.assertIsNone(phones.total)
        self.assertEqual(phones.count, 2)
        self.assertNotIn("交教会日期", {item.name for item in by_meet.header_stats})

    def test_analyze_splits_channel_header_into_series(self):
        records = [
            Record("s", "交教会", "g", "浇灌", 2, {
                "交教会日期": "2026-10-05", "加友途径4": "Facebook", "摸底/推广": "张三",
            }, "a"),
            Record("s", "交教会", "g", "浇灌", 3, {
                "交教会日期": "2026-10-05", "加友途径4": "WhatsApp", "摸底/推广": "李四",
            }, "b"),
            Record("s", "交教会", "g", "浇灌", 4, {
                "交教会日期": "2026-10-06", "加友途径4": "Facebook", "摸底/推广": "张三",
            }, "c"),
            Record("s", "交教会", "g", "浇灌", 5, {
                "交教会日期": "2026-09-28", "加友途径4": "Facebook", "摸底/推广": "张三",
            }, "d"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(records)
            compared = engine.analyze(
                source="aggregate",
                date_field="交教会日期",
                compare_mode="week",
                compare=True,
                reference_date=date(2026, 10, 6),
            )
            current_only = engine.analyze(
                source="aggregate",
                date_field="交教会日期",
                compare_mode="month",
                compare=False,
                reference_date=date(2026, 10, 6),
            )
        channels = {item.header: item for item in compared.breakdowns}["加友途径4"]
        series = {part.label: part for part in channels.series}
        self.assertEqual(series["Facebook"].count, 2)
        self.assertEqual(series["WhatsApp"].count, 1)
        self.assertEqual(series["Facebook"].previous_count, 1)
        self.assertEqual(sum(series["Facebook"].daily), 2)
        self.assertTrue(compared.compare_enabled)
        self.assertFalse(current_only.compare_enabled)
        self.assertEqual(current_only.previous.count, 0)
        self.assertEqual(current_only.current.count, 3)
        self.assertEqual(current_only.current.start, "2026-10-01")

    def test_analyze_uses_selected_name_column(self):
        records = [
            Record("s", "交教会", "g", "浇灌", 2, {
                "交教会日期": "2026-10-05", "摸底/推广": "张小川", "线索名字": "小明",
            }, "a"),
            Record("s", "交教会", "g", "浇灌", 3, {
                "交教会日期": "2026-10-06", "摸底/推广": "张小川", "线索名字": "小红",
            }, "b"),
            Record("s", "交教会", "g", "浇灌", 4, {
                "交教会日期": "2026-10-06", "摸底/推广": "李四", "线索名字": "小明",
            }, "c"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(records)
            by_lead = engine.analyze(
                source="aggregate",
                date_field="交教会日期",
                name_field="线索名字",
                names=["小明"],
                compare_mode="week",
                compare=False,
                reference_date=date(2026, 10, 6),
            )
            by_staff = engine.analyze(
                source="aggregate",
                date_field="交教会日期",
                name_field="摸底/推广",
                names=["张小川"],
                compare_mode="week",
                compare=False,
                reference_date=date(2026, 10, 6),
            )
        self.assertEqual(by_lead.name_field, "线索名字")
        self.assertEqual(by_lead.current.count, 2)
        self.assertEqual(by_staff.name_field, "摸底/推广")
        self.assertEqual(by_staff.current.count, 2)

    def test_analyze_session_header_counts_d_not_durations(self):
        records = [
            Record("s", "交教会", "g", "浇灌", 2, {
                "交教会日期": "2026-10-05", "第一场": "Não", "摸底/推广": "张三",
            }, "a"),
            Record("s", "交教会", "g", "浇灌", 3, {
                "交教会日期": "2026-10-05", "第一场": "10.6 D 97min/笔聊73min", "摸底/推广": "李四",
            }, "b"),
            Record("s", "交教会", "g", "浇灌", 4, {
                "交教会日期": "2026-10-06", "第一场": "10.2 1min", "摸底/推广": "张三",
            }, "c"),
            Record("s", "交教会", "g", "浇灌", 5, {
                "交教会日期": "2026-10-06", "第一场": "①D 2min", "摸底/推广": "王五",
            }, "d"),
        ]
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all(records)
            result = engine.analyze(
                source="aggregate",
                date_field="交教会日期",
                compare_mode="week",
                compare=False,
                reference_date=date(2026, 10, 6),
            )
        sessions = {item.header: item for item in result.breakdowns}["第一场"]
        self.assertEqual(len(sessions.series), 1)
        self.assertEqual(sessions.series[0].label, "第一场")
        self.assertEqual(sessions.series[0].count, 2)
        daily = {point.day: sessions.series[0].daily[index] for index, point in enumerate(result.current.daily)}
        self.assertEqual(daily["2026-10-05"], 1)
        self.assertEqual(daily["2026-10-06"], 1)
        self.assertNotIn("Não", {part.label for part in sessions.series})

    def test_source_cache_merge_skips_unchanged_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            cache = SourceCache(Path(directory))
            first_rows = [Record("s", "源", "g", "表", 2, {"名字": "张三"}, "h1")]
            first = cache.merge("s", first_rows, ["名字"])
            self.assertEqual(first["inserted"], 1)
            self.assertFalse(first["unchanged"])
            same = cache.merge("s", first_rows, ["名字"])
            self.assertTrue(same["unchanged"])
            self.assertEqual(same["skipped"], 1)
            self.assertEqual(same["updated"], 0)
            changed = [Record("s", "源", "g", "表", 2, {"名字": "李四"}, "h2")]
            updated = cache.merge("s", changed, ["名字"])
            self.assertEqual(updated["updated"], 1)
            self.assertEqual(updated["inserted"], 0)
            self.assertFalse(updated["unchanged"])
            self.assertEqual(cache.load("s")[0].values["名字"], "李四")

    def test_cached_records_require_local_without_sync(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            with self.assertRaises(ValueError) as ctx:
                engine.cached_records(["missing-id"], require_local=True)
            self.assertIn("同步本地库", str(ctx.exception))

    def test_batch_query_exact_uses_index_and_keeps_all_matches(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "A", 2, {"号码": "258-851-758692"}, "h1"),
                Record("s", "源", "g", "B", 3, {"手机号码": "258851758692"}, "h2"),
                Record("s", "源", "g", "C", 4, {"号码": "999"}, "h3"),
            ])
            results = engine.query_many("号码", ["258851758692", "000"], exact=True)
            self.assertEqual([record.sheet_name if record else None for _, record in results], ["A", "B", None])

    def test_batch_query_fuzzy_reuses_prepared_values(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1008-李薇", 2, {"号码": "123"}, "h1"),
                Record("s", "源", "g", "AAOZ-依心", 3, {"号码": "456"}, "h2"),
            ])
            results = engine.query_many("来源", ["李薇", "依心"], exact=False)
            self.assertEqual([record.sheet_name for _, record in results if record], ["1008-李薇", "AAOZ-依心"])

    def test_fuzzy_source_query_matches_original_and_corrected_sheet_name(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1233-赵刚", 2, {"号码": "123"}, "h1")
            ])
            for value in ("赵刚", "1233", "1233-赵刚", "赵刚-1233-专页后台"):
                results = engine.query_many("来源", [value], exact=False)
                self.assertEqual(results[0][1].sheet_name, "1233-赵刚")

    def test_exact_source_query_matches_corrected_sheet_name(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "1233-赵刚", 2, {"号码": "123"}, "h1")
            ])
            results = engine.query_many("来源", ["赵刚-1233-专页后台"], exact=True)
            self.assertEqual(results[0][1].sheet_name, "1233-赵刚")

    def test_query_can_limit_results_by_date_range(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record("s", "源", "g", "A", 2, {"名字": "刘海", "日期": "2026-09-01"}, "h1"),
                Record("s", "源", "g", "B", 3, {"名字": "刘海", "日期": "2026-09-13"}, "h2"),
                Record("s", "源", "g", "C", 4, {"名字": "刘海", "日期": "无效日期"}, "h3"),
            ])
            results = engine.query_many(
                "名字", ["刘海"], exact=True, date_field="日期",
                start_date=date(2026, 9, 10), end_date=date(2026, 9, 20),
            )
            self.assertEqual(len(results), 1)
            self.assertEqual(results[0][1].sheet_name, "B")

    def test_custom_headers_are_resolved_for_query_date_and_phone(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "A", 2,
                    {
                        "交教会日期": "2026年9月13日",
                        "线索电话号码": "258-851-758692",
                        "摸底/推广": "简\u200b 单",
                    },
                    "h1",
                ),
            ])
            results = engine.query_many(
                "摸底/推广", ["简单"], exact=False, date_field="日期",
                start_date=date(2026, 9, 1), end_date=date(2026, 9, 30),
            )
            self.assertIsNotNone(results[0][1])
            self.assertEqual(engine.query("号码", "258851758692")[0].sheet_name, "A")
            self.assertEqual(
                engine.suggest_field(["交教会日期", "摸底/推广"], "日期"),
                "交教会日期",
            )

    def test_extract_resolves_custom_date_and_dedup_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "A", 2,
                    {"交教会日期": "2026-09-13", "线索电话号码": "123"},
                    "h1",
                ),
            ])
            first = engine.extract(
                "日期", date(2026, 9, 1), date(2026, 9, 30),
                root / "one.xlsx", ["号码", "日期"],
            )
            second = engine.extract(
                "日期", date(2026, 9, 1), date(2026, 9, 30),
                root / "two.xlsx", ["号码", "日期"],
            )
            self.assertEqual(first["written"], 1)
            self.assertEqual(second["duplicates"], 1)

    def test_extract_dedup_reads_existing_output_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "existing.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "提取结果"
            sheet.append(["来源", "手机号码", "日期"])
            sheet.append(["历史来源", "258-851-758692", "2026-09-13"])
            workbook.save(output)
            workbook.close()

            store = ConfigStore(root / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "新来源", 2,
                    {"号码": "258851758692", "日期": "2026-09-13"},
                    "h1",
                ),
            ])
            result = engine.extract(
                "日期", date(2026, 9, 1), date(2026, 9, 30),
                output, ["号码", "日期"], output_sheet_name="提取结果",
            )
            self.assertEqual(result["written"], 0)
            self.assertEqual(result["duplicates"], 1)
            reopened = openpyxl.load_workbook(output, data_only=True)
            try:
                rows = list(reopened["提取结果"].iter_rows(values_only=True))
            finally:
                reopened.close()
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[1][0], "历史来源")

    def test_extract_inserts_new_rows_at_top_of_existing_xlsx(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "existing.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "提取结果"
            sheet.append(["来源", "号码", "日期"])
            sheet.append(["旧来源", "111", "2026-09-01"])
            workbook.save(output)
            workbook.close()

            store = ConfigStore(root / "config")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "新来源", 2,
                    {"号码": "222", "日期": "2026-09-13"},
                    "h1",
                ),
            ])
            result = engine.extract(
                "日期", date(2026, 9, 1), date(2026, 9, 30),
                output, ["号码", "日期"], output_sheet_name="提取结果",
            )
            self.assertEqual(result["written"], 1)
            reopened = openpyxl.load_workbook(output, data_only=True)
            try:
                rows = list(reopened["提取结果"].iter_rows(values_only=True))
            finally:
                reopened.close()
            self.assertEqual(rows[1], ("新来源", "222", "2026-09-13"))
            self.assertEqual(rows[2], ("旧来源", "111", "2026-09-01"))

    def test_direct_extract_can_limit_source_and_adds_source_header(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            store.set("column_schema_enabled", True)
            store.set("column_schema", ["专页ID", "号码"])
            engine = DataEngine(store)
            records = [
                Record(
                    "src2", "源二", "g", "浇灌数据库-过滤", 2,
                    {
                        "交教会日期": "2026-09-13",
                        "线索电话号码": "258851758692",
                        "摸底/推广": "简单",
                    },
                    "h1",
                )
            ]
            output = root / "direct_extract.xlsx"
            with patch.object(engine, "read_sources", return_value=records) as read_sources:
                result = engine.extract(
                    "日期", date(2026, 9, 1), date(2026, 9, 30),
                    output, ["号码", "日期"], direct=True, source_id="src2",
                )
            read_sources.assert_called_once_with("时间提取", "src2")
            self.assertEqual(result["written"], 1)
            workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows = list(sheet.iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(rows[0], ("来源", "交教会日期", "线索电话号码", "摸底/推广"))
            self.assertEqual(rows[1], ("浇灌数据库-过滤", "2026-09-13", "258851758692", "简单"))

    def test_direct_extract_signature_column_is_added_to_new_output(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            store = ConfigStore(root / "config")
            store.set("extract_signature_enabled", True)
            store.set("extract_signature_header", "签字")
            store.set("extract_signature_value", "张三")
            store.set("extract_signature_column", "K")
            engine = DataEngine(store)
            output = root / "signature.xlsx"
            records = [
                Record(
                    "src", "源", "g", "浇灌数据库-过滤", 2,
                    {"交教会日期": "2026-09-13", "线索电话号码": "258851758692", "摸底/推广": "简单"},
                    "h1",
                )
            ]
            with patch.object(engine, "read_sources", return_value=records):
                result = engine.extract(
                    "日期", date(2026, 9, 1), date(2026, 9, 30),
                    output, ["号码", "日期"], direct=True,
                )
            self.assertEqual(result["written"], 1)
            workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual(rows[0][0], "来源")
            self.assertEqual(rows[0][10], "签字")
            self.assertEqual(rows[1][10], "张三")

    def test_extract_output_schema_supports_column_mapping_objects(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("extract_column_schema_enabled", True)
            store.set("extract_column_schema", [
                {"name": "线索电话号码", "column": "S", "enabled": True},
                {"name": "交教会日期", "column": "F", "enabled": True},
            ])
            engine = DataEngine(store)
            headers = engine._preferred_export_headers([])
            self.assertEqual(headers, ["来源", "线索电话号码", "交教会日期"])
            rows = engine._rows_for_headers(
                [Record("s", "源", "g", "A", 2, {"号码": "123", "日期": "2026-09-13"}, "h")],
                headers,
                store.get("field_aliases"),
            )
            self.assertEqual(rows[0], ["A", "123", "2026-09-13"])

    def test_query_extract_table_is_faster_subset(self):
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            engine = DataEngine(store)
            extract_path = Path(directory) / "extract.xlsx"
            DataEngine._write_xlsx(
                extract_path,
                [
                    Record("s", "源", "g", "1008-李薇", 2, {"号码": "123", "名字": "李薇", "日期": "2026-09-13"}, "h1"),
                    Record("s", "源", "g", "2002-王强", 3, {"号码": "456", "名字": "王强", "日期": "2026-09-12"}, "h2"),
                ],
                "提取结果",
            )
            other = openpyxl.load_workbook(extract_path)
            extra = other.create_sheet("其他页")
            extra.append(["来源", "号码"])
            extra.append(["不该读", "999"])
            other.save(extract_path)
            other.close()
            results = engine.query_many(
                "号码",
                ["123", "999"],
                source="extract",
                extract_target=str(extract_path),
                extract_sheet="提取结果",
                refresh_cache=True,
            )
            self.assertEqual(results[0][1].sheet_name, "1008-李薇")
            self.assertEqual(results[0][1].values["号码"], "123")
            self.assertIsNone(results[1][1])
            found = engine.query(
                "号码",
                "456",
                source="extract",
                extract_target=str(extract_path),
                extract_sheet="提取结果",
            )
            self.assertEqual(found[0].values["名字"], "王强")
            headers = engine.list_query_fields(
                "extract", str(extract_path), "提取结果",
            )
            self.assertIn("来源", headers)
            self.assertTrue(any(name in {"号码", "手机号码"} for name in headers))
            by_header = engine.query(
                "名字",
                "李薇",
                source="extract",
                extract_target=str(extract_path),
                extract_sheet="提取结果",
            )
            self.assertEqual(by_header[0].values["号码"], "123")

    def test_extract_uses_configured_sheet_and_source_first(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.xlsx"
            records = [Record("s", "源", "g", "渠道Sheet", 2, {"号码": "123", "日期": "2026-09-13"}, "h")]
            DataEngine._write_xlsx(path, records, "指定工作表")
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self.assertEqual(book.sheetnames, ["指定工作表"])
            rows = list(book["指定工作表"].iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "来源")
            self.assertEqual(rows[1][0], "渠道Sheet")
            book.close()

    def test_export_header_aliases_and_order_are_compatible(self):
        aliases = {
            "名字": ["姓名"],
            "号码": ["手机号码"],
            "日期": ["日期"],
        }
        existing = ["来源", "专页ID", "姓名", "手机号码", "日期"]
        required = ["来源", "日期", "名字", "号码", "专页ID"]
        self.assertTrue(DataEngine._headers_compatible(existing, required, aliases))
        records = [Record(
            "s", "源", "g", "1008-李薇", 2,
            {"专页ID": "42", "名字": "李薇", "号码": "258851758692", "日期": "2026-09-13"},
            "h",
        )]
        self.assertEqual(
            DataEngine._rows_for_headers(records, existing, aliases)[0],
            ["1008-李薇", "42", "李薇", "258851758692", "2026-09-13"],
        )


class WorkbookTests(unittest.TestCase):
    def test_local_workbook_source(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "订单"
            sheet.append(["日期", "手机号", "渠道"])
            sheet.append(["2026-09-13", "10086", "搜索"])
            stats = book.create_sheet("统计")
            stats.append(["日期", "手机号"])
            stats.append(["2026-09-13", "should_skip"])
            book.save(path)
            reader = SourceReader({"号码": ["手机号"], "来源": ["渠道"], "日期": ["日期"]}, ["统计"])
            source = SourceConfig("id", "测试", str(path))
            records = reader.read(source)
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0].values["号码"], "10086")

    def test_sources_can_use_different_column_counts(self):
        with tempfile.TemporaryDirectory() as directory:
            short_path = Path(directory) / "short.xlsx"
            long_path = Path(directory) / "long.xlsx"
            short = openpyxl.Workbook()
            short.active.title = "短表"
            short.active.append(["a", "b", "ignore"])
            short.active.append(["page-1", "13800138000", "skip"])
            short.save(short_path)
            long = openpyxl.Workbook()
            long.active.title = "长表"
            long.active.append(["a", "b", "c", "d"])
            long.active.append(["page-2", "李薇", "https://example.test", "258851758692"])
            long.save(long_path)
            store = ConfigStore(Path(directory) / "config")
            store.save_source(SourceConfig(
                "short", "短表", str(short_path),
                column_schema_enabled=True,
                column_schema=["专页ID", "手机号码"],
            ))
            store.save_source(SourceConfig(
                "long", "长表", str(long_path),
                column_schema_enabled=True,
                column_schema=["专页ID", "姓名", "评论贴文", "手机号码"],
            ))
            engine = DataEngine(store)
            records = engine.read_sources()
            by_source = {record.source_id: record for record in records}
            self.assertEqual(by_source["short"].values["专页ID"], "page-1")
            self.assertEqual(by_source["short"].values["号码"], "13800138000")
            self.assertNotIn("ignore", by_source["short"].values)
            self.assertEqual(by_source["long"].values["名字"], "李薇")
            self.assertEqual(by_source["long"].values["号码"], "258851758692")
            self.assertEqual(by_source["long"].values["评论贴文"], "https://example.test")
            short_fields = engine.list_query_fields("direct", source_id="short")
            long_fields = engine.list_query_fields("direct", source_id="long")
            self.assertIn("手机号码", short_fields)
            self.assertNotIn("评论贴文", short_fields)
            self.assertIn("评论贴文", long_fields)
            self.assertEqual(
                engine.query(
                    "手机号码", "13800138000", source="direct", source_id="short", refresh_cache=True,
                )[0].source_id,
                "short",
            )

    def test_manual_column_count_and_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manual.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "数据"
            sheet.append(["原表头A", "原表头B", "不读取的列"])
            sheet.append(["page-1", "13800138000", "ignore-me"])
            book.save(path)
            reader = SourceReader(
                {"号码": ["手机号码"]},
                [],
                column_schema=["专页ID", "手机号码"],
            )
            records = reader.read(SourceConfig("id", "手动列", str(path)))
            self.assertEqual(records[0].values["专页ID"], "page-1")
            self.assertEqual(records[0].values["号码"], "13800138000")
            self.assertNotIn("不读取的列", records[0].values)

    def test_column_letter_mapping_skips_leading_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "letters.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "数据"
            sheet.append(["A列", "B列", "C列", "贴文ID", "手机号码", "多余"])
            sheet.append(["x", "y", "z", "page-9", "13800138000", "no"])
            book.save(path)
            reader = SourceReader(
                {"号码": ["手机号码"]},
                [],
                column_schema=[
                    {"name": "贴文ID", "column": "D", "enabled": True},
                    {"name": "手机号码", "column": "E", "enabled": True},
                    {"name": "忽略", "column": "F", "enabled": False},
                ],
            )
            records = reader.read(SourceConfig("id", "列映射", str(path)))
            self.assertEqual(records[0].values["贴文ID"], "page-9")
            self.assertEqual(records[0].values["号码"], "13800138000")
            self.assertNotIn("忽略", records[0].values)
            self.assertNotIn("A列", records[0].values)

    def test_private_source_uses_one_batch_request(self):
        class FakeSpreadsheet:
            def __init__(self):
                self.batch_calls = 0

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [
                    {"properties": {"title": "A"}},
                    {"properties": {"title": "B"}},
                ]}

            def values_batch_get(self, key, ranges, params=None):
                self.batch_calls += 1
                return {"valueRanges": [
                    {"values": [["手机号"], ["111"]]},
                    {"values": [["手机号"], ["222"]]},
                ]}

        spreadsheet = FakeSpreadsheet()
        client = SimpleNamespace(http_client=spreadsheet)
        source = SourceConfig("id", "私有表", "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz", credential_path="fake.json")
        reader = SourceReader({"号码": ["手机号"]}, [])
        with patch.object(SourceReader, "_gspread_client", return_value=client):
            records = reader.read(source)
        self.assertEqual(spreadsheet.batch_calls, 1)
        self.assertEqual([record.values["号码"] for record in records], ["111", "222"])

    def test_google_output_appends_header_and_rows(self):
        class FakeHttp:
            def __init__(self):
                self.inserted = []
                self.written = []

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [{"properties": {"sheetId": 1, "title": "目标页"}}]}

            def values_get(self, key, range_name, params=None):
                return {"values": []}

            def batch_update(self, key, body=None):
                self.inserted.extend(body["requests"])

            def values_update(self, key, range_name, params=None, body=None):
                self.written.extend(body["values"])

        http = FakeHttp()
        client = SimpleNamespace(http_client=http)
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", "fake.json")
            engine = DataEngine(store)
            records = [Record("s", "源", "g", "渠道A", 2, {"号码": "123"}, "h")]
            with patch.object(SourceReader, "_gspread_client", return_value=client):
                engine._write_google_sheet(
                    "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz",
                    "目标页",
                    records,
                    "测试",
                )
        self.assertEqual(http.written[0], ["来源", "号码"])
        self.assertEqual(http.written[1], ["渠道A", "123"])

    def test_google_extract_dedup_reads_existing_target_sheet(self):
        existing = ["来源", "手机号码", "日期"]

        class FakeHttp:
            def __init__(self):
                self.inserted = []
                self.written = []

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [{"properties": {"sheetId": 1, "title": "提取结果"}}]}

            def values_get(self, key, range_name, params=None):
                if str(range_name).endswith("!A:ZZZ"):
                    return {"values": [existing, ["旧来源", "258-851-758692", "2026-09-13"]]}
                return {"values": [existing]}

            def batch_update(self, key, body=None):
                self.inserted.extend(body["requests"])

            def values_update(self, key, range_name, params=None, body=None):
                self.written.extend(body["values"])

        http = FakeHttp()
        client = SimpleNamespace(http_client=http)
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", "fake.json")
            engine = DataEngine(store)
            engine.database.replace_all([
                Record(
                    "s", "源", "g", "新来源", 2,
                    {"号码": "258851758692", "日期": "2026-09-13"},
                    "h1",
                ),
            ])
            with patch.object(SourceReader, "_gspread_client", return_value=client):
                result = engine.extract(
                    "日期", date(2026, 9, 1), date(2026, 9, 30),
                    Path(directory) / "unused.xlsx", ["号码", "日期"],
                    destination_type="google",
                    google_output_url="https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz",
                    output_sheet_name="提取结果",
                )
        self.assertEqual(result["written"], 0)
        self.assertEqual(result["duplicates"], 1)
        self.assertEqual(http.inserted, [])
        self.assertEqual(http.written, [])

    def test_google_output_uses_existing_alias_header_order(self):
        existing = ["来源", "专页ID", "姓名", "标签", "订阅时间", "性别", "评论贴文", "手机号码", "日期"]

        class FakeHttp:
            def __init__(self):
                self.inserted = []
                self.written = []

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [{"properties": {"sheetId": 1, "title": "测试"}}]}

            def values_get(self, key, range_name, params=None):
                return {"values": [existing]}

            def batch_update(self, key, body=None):
                self.inserted.extend(body["requests"])

            def values_update(self, key, range_name, params=None, body=None):
                self.written.extend(body["values"])

        http = FakeHttp()
        client = SimpleNamespace(http_client=http)
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", "fake.json")
            store.set("column_schema_enabled", True)
            store.set("column_schema", existing[1:])
            engine = DataEngine(store)
            records = [Record(
                "s", "源", "g", "1008-李薇", 2,
                {
                    "专页ID": "42", "名字": "李薇", "标签": "A", "订阅时间": "12:00",
                    "性别": "女", "评论贴文": "https://example.test", "号码": "258851758692",
                    "日期": "2026-09-13",
                },
                "h",
            )]
            with patch.object(SourceReader, "_gspread_client", return_value=client):
                engine._write_google_sheet(
                    "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz",
                    "测试",
                    records,
                    "测试",
                )
        self.assertEqual(http.inserted[0]["insertDimension"]["range"]["startIndex"], 1)
        self.assertEqual(http.written[0], [
            "1008-李薇", "42", "李薇", "A", "12:00", "女", "https://example.test",
            "258851758692", "2026-09-13",
        ])

    def test_google_output_uses_existing_headers_without_rejecting_mismatch(self):
        existing = ["来源", "专页ID", "姓名", "标签", "订阅时间", "性别", "评论贴文", "手机号码", "日期", "签字"]

        class FakeHttp:
            def __init__(self):
                self.inserted = []
                self.written = []

            def fetch_sheet_metadata(self, key, params=None):
                return {"sheets": [{"properties": {"sheetId": 1, "title": "提取表格"}}]}

            def values_get(self, key, range_name, params=None):
                return {"values": [existing]}

            def batch_update(self, key, body=None):
                self.inserted.extend(body["requests"])

            def values_update(self, key, range_name, params=None, body=None):
                self.written.extend(body["values"])

        http = FakeHttp()
        client = SimpleNamespace(http_client=http)
        with tempfile.TemporaryDirectory() as directory:
            store = ConfigStore(Path(directory) / "config")
            store.set("credential_path", "fake.json")
            store.set("extract_signature_enabled", True)
            store.set("extract_signature_header", "签字")
            store.set("extract_signature_value", "张三")
            engine = DataEngine(store)
            records = [Record(
                "s", "源", "g", "1008-李薇", 2,
                {
                    "日期": "2026-09-13", "名字": "李薇", "号码": "258851758692",
                    "专页ID": "42", "评论贴文": "https://example.test",
                },
                "h",
            )]
            with patch.object(SourceReader, "_gspread_client", return_value=client):
                engine._write_google_sheet(
                    "https://docs.google.com/spreadsheets/d/abcdefghijklmnopqrstuvwxyz",
                    "提取表格",
                    records,
                    "测试",
                    prefer_record_headers=True,
                )
        self.assertEqual(http.inserted[0]["insertDimension"]["range"]["startIndex"], 1)
        self.assertEqual(http.written[0], [
            "1008-李薇", "42", "李薇", "", "", "", "https://example.test",
            "258851758692", "2026-09-13", "张三",
        ])


if __name__ == "__main__":
    unittest.main()
