from __future__ import annotations

import hashlib
import io
import json
import random
import re
import time
from pathlib import Path
from typing import Callable, Iterable
from urllib.parse import urlparse

import openpyxl
import requests

from .models import Record, SourceConfig


LogFn = Callable[[str, str], None]
MAX_XLSX_BYTES = 80 * 1024 * 1024
XLSX_MAGIC = b"PK"
GOOGLE_DOWNLOAD_HOSTS = {
    "docs.google.com",
    "drive.google.com",
    "spreadsheets.google.com",
    "googleusercontent.com",
}


def _google_download_host_ok(url: str) -> bool:
    host = (urlparse(str(url or "")).hostname or "").casefold()
    if not host:
        return False
    if host in GOOGLE_DOWNLOAD_HOSTS:
        return True
    return host.endswith(".google.com") or host.endswith(".googleusercontent.com")


RETRY_STATUSES = {429, 500, 502, 503, 504}


def normalize_credential_paths(*groups: Iterable[object]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for group in groups:
        if group is None:
            continue
        items = [group] if isinstance(group, str) else group
        for item in items:
            text = str(item or "").strip()
            key = text.casefold()
            if text and key not in seen:
                seen.add(key)
                result.append(text)
    return result


def google_retry(call, logger: LogFn | None = None, attempts: int = 7):
    """Retry Google API quota and transient server errors with bounded backoff."""
    log = logger or (lambda level, message: None)
    for attempt in range(attempts):
        try:
            return call()
        except Exception as exc:
            response = getattr(exc, "response", None)
            status = getattr(response, "status_code", None)
            if status not in RETRY_STATUSES or attempt == attempts - 1:
                raise
            wait_seconds = min(45.0, (2 ** attempt) + random.random())
            log("WARNING", f"Google API 暂时限流（{status}），{wait_seconds:.1f} 秒后自动重试 {attempt + 2}/{attempts}")
            time.sleep(wait_seconds)


class CredentialPool:
    """Round-robin Google service accounts; rotate immediately on 429/5xx."""

    def __init__(
        self,
        paths: Iterable[str] | None,
        logger: LogFn | None = None,
        client_factory=None,
        start: int = 0,
    ):
        self.paths = normalize_credential_paths(paths)
        self.logger = logger or (lambda level, message: None)
        self._factory = client_factory
        self._clients: dict[str, object] = {}
        self.index = int(start) % len(self.paths) if self.paths else 0

    def _factory_fn(self):
        return self._factory or SourceReader._gspread_client

    def _client(self, path: str):
        if path not in self._clients:
            self._clients[path] = self._factory_fn()(path)
        return self._clients[path]

    def call(self, fn, attempts: int = 7):
        if not self.paths:
            raise FileNotFoundError("没有可用的服务账号 JSON")
        last_exc: Exception | None = None
        retries = max(int(attempts), len(self.paths))
        for attempt in range(retries):
            path = self.paths[self.index % len(self.paths)]
            try:
                result = fn(self._client(path))
                self.index = (self.index + 1) % len(self.paths)
                return result
            except FileNotFoundError as exc:
                last_exc = exc
                self.logger("WARNING", f"服务账号无效，已跳过：{Path(path).name}")
                self.index = (self.index + 1) % len(self.paths)
                if attempt == retries - 1:
                    raise
            except Exception as exc:
                last_exc = exc
                status = getattr(getattr(exc, "response", None), "status_code", None)
                if status in RETRY_STATUSES and attempt < retries - 1:
                    self.logger("WARNING", f"服务账号 {Path(path).name} 返回 {status}，切换下一个")
                    self.index = (self.index + 1) % len(self.paths)
                    if (attempt + 1) % len(self.paths) == 0:
                        wait_seconds = min(20.0, (1.6 ** (attempt // max(1, len(self.paths)))) + random.random())
                        self.logger("WARNING", f"全部服务账号均限流，{wait_seconds:.1f} 秒后继续轮询")
                        time.sleep(wait_seconds)
                    continue
                raise
        if last_exc:
            raise last_exc
        raise RuntimeError("服务账号轮询失败")


def spreadsheet_id(url: str) -> str:
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9_-]+)", url)
    if match:
        return match.group(1)
    if re.fullmatch(r"[a-zA-Z0-9_-]{20,}", url.strip()):
        return url.strip()
    return ""


def split_names(text: str | Iterable[str]) -> list[str]:
    if not isinstance(text, str):
        return [str(item).strip() for item in text if str(item).strip()]
    return [item.strip() for item in re.split(r"[,，;；\n]+", text) if item.strip()]


def parse_schema_lines(text: str | Iterable[str]) -> list[str]:
    if not isinstance(text, str):
        return [str(item).strip() for item in text]
    lines = [line.strip() for line in text.splitlines()]
    while lines and not lines[-1]:
        lines.pop()
    return lines


def excel_column(index: int) -> str:
    result = ""
    number = max(0, int(index)) + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(65 + remainder) + result
    return result or "A"


def column_index(letter: str) -> int:
    text = re.sub(r"[^A-Za-z]", "", str(letter or "")).upper()
    if not text:
        return 0
    value = 0
    for char in text:
        value = value * 26 + (ord(char) - 64)
    return max(0, value - 1)


def normalize_column_schema(schema: Iterable[object] | None) -> list[dict[str, object]]:
    result: list[dict[str, object]] = []
    for index, item in enumerate(schema or []):
        if isinstance(item, str):
            name = item.strip()
            result.append({"name": name, "column": excel_column(index), "enabled": bool(name)})
            continue
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or item.get("header") or "").strip()
        column = str(item.get("column") or excel_column(index)).strip().upper() or excel_column(index)
        enabled = bool(item.get("enabled", True))
        result.append({"name": name, "column": column, "enabled": enabled})
    return result


def schema_field_names(schema: Iterable[object] | None) -> list[str]:
    return [str(item["name"]).strip() for item in normalize_column_schema(schema) if item.get("enabled") and str(item.get("name") or "").strip()]


def choose_sheets(all_names: list[str], include: list[str], exclude: list[str]) -> tuple[list[str], list[str]]:
    available_names = {name.casefold() for name in all_names}

    def expand(configured: list[str]) -> list[str]:
        expanded: list[str] = []
        for name in configured:
            clean = name.strip()
            if clean.casefold() in available_names:
                expanded.append(clean)
            else:
                expanded.extend(part for part in re.split(r"\s+", clean) if part)
        return expanded

    include = expand(include)
    exclude = expand(exclude)
    excluded = {name.casefold() for name in exclude}
    available = {name.casefold(): name for name in all_names}
    requested = include if include else all_names
    selected: list[str] = []
    missing: list[str] = []
    for name in requested:
        actual = available.get(name.casefold())
        if actual is None:
            missing.append(name)
        elif actual.casefold() not in excluded and actual not in selected:
            selected.append(actual)
    return selected, missing


def normalize_headers(values: list[object]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"未命名列{index}"
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def canonicalize(row: dict[str, object], aliases: dict[str, list[str]]) -> dict[str, str]:
    lookup = {str(key).strip().casefold(): value for key, value in row.items()}
    result: dict[str, str] = {}
    used: set[str] = set()
    for canonical, candidates in aliases.items():
        for candidate in [canonical, *candidates]:
            key = str(candidate).strip().casefold()
            if key in lookup:
                value = lookup[key]
                result[canonical] = "" if value is None else str(value).strip()
                used.add(key)
                break
    for header, value in row.items():
        if str(header).strip().casefold() not in used:
            result[str(header)] = "" if value is None else str(value).strip()
    return result


class SourceReader:
    def __init__(
        self,
        aliases: dict[str, list[str]],
        global_excludes: list[str],
        logger: LogFn | None = None,
        column_schema: list[str] | None = None,
        credential_paths: list[str] | None = None,
    ):
        self.aliases = aliases
        self.global_excludes = global_excludes
        self.logger = logger or (lambda level, message: None)
        self.column_schema = list(column_schema or [])
        self.credential_paths = normalize_credential_paths(credential_paths)

    def _headers(self, actual_values: list[object]) -> tuple[list[str], list[int] | None]:
        entries = normalize_column_schema(self.column_schema)
        if not entries:
            return normalize_headers(actual_values), None
        names: list[str] = []
        indices: list[int] = []
        padded = list(actual_values)
        for entry in entries:
            if not entry.get("enabled", True):
                continue
            index = column_index(str(entry.get("column") or "A"))
            original = str(padded[index]).strip() if 0 <= index < len(padded) else ""
            name = str(entry.get("name") or "").strip() or original
            if not name:
                continue
            names.append(name)
            indices.append(index)
        if not names:
            return normalize_headers(actual_values), None
        return normalize_headers(names), indices

    @staticmethod
    def _cells(values: list[object], headers: list[str], selectors: list[int] | None) -> dict[str, object]:
        cells = list(values)
        if selectors is None:
            limited = cells[:len(headers)] if headers else cells
            return dict(zip(headers, limited))
        mapped: dict[str, object] = {}
        for name, index in zip(headers, selectors):
            mapped[name] = cells[index] if 0 <= index < len(cells) else ""
        return mapped

    def read(self, source: SourceConfig) -> list[Record]:
        sid = spreadsheet_id(source.url)
        if not sid:
            path = Path(source.url).expanduser()
            if path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}:
                content = path.read_bytes()
                if len(content) > MAX_XLSX_BYTES:
                    raise ValueError("本地 Excel 文件过大，已拒绝读取")
                return self._read_workbook(source, content, path.stem)
            raise ValueError("无法识别 Google 表格链接或本地 Excel 文件")
        paths = normalize_credential_paths(source.credential_path, self.credential_paths)
        if paths:
            return self._read_private(source, sid, paths)
        return self._read_public(source, sid)

    def list_sheets(self, source: SourceConfig) -> list[str]:
        sid = spreadsheet_id(source.url)
        paths = normalize_credential_paths(source.credential_path, self.credential_paths)
        if paths:
            pool = CredentialPool(paths, self.logger)
            metadata = pool.call(
                lambda client: client.http_client.fetch_sheet_metadata(
                    sid, params={"fields": "sheets.properties.title"}
                )
            )
            return [item["properties"]["title"] for item in metadata.get("sheets", [])]
        content = self._download_public_xlsx(sid, timeout=60)
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        names = workbook.sheetnames
        workbook.close()
        return names

    def _read_public(self, source: SourceConfig, sid: str) -> list[Record]:
        self.logger("INFO", f"下载公开表格：{source.name}")
        content = self._download_public_xlsx(sid, timeout=120)
        return self._read_workbook(source, content, sid)

    @staticmethod
    def _download_public_xlsx(sid: str, timeout: int = 120) -> bytes:
        if not re.fullmatch(r"[a-zA-Z0-9_-]{20,}", sid or ""):
            raise ValueError("表格 ID 无效")
        response = requests.get(
            f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx",
            timeout=timeout,
            stream=True,
            headers={"User-Agent": "SheetDataHub"},
        )
        if response.status_code in {401, 403}:
            raise PermissionError("表格不可公开读取，请共享为可查看或配置服务账号 JSON")
        response.raise_for_status()
        if not _google_download_host_ok(str(getattr(response, "url", "") or "")):
            raise ValueError("公开表格下载被重定向到非 Google 地址，已拒绝")
        chunks: list[bytes] = []
        total = 0
        for chunk in response.iter_content(chunk_size=1024 * 256):
            if not chunk:
                continue
            total += len(chunk)
            if total > MAX_XLSX_BYTES:
                raise ValueError("公开表格下载超过 80MB，已中止")
            chunks.append(chunk)
        content = b"".join(chunks)
        if not content.startswith(XLSX_MAGIC):
            raise ValueError("下载内容不是有效的 Excel 文件")
        return content

    def _read_workbook(self, source: SourceConfig, content: bytes, sid: str) -> list[Record]:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        excludes = [*self.global_excludes, *source.exclude_sheets]
        selected, missing = choose_sheets(workbook.sheetnames, source.include_sheets, excludes)
        for name in missing:
            self.logger("WARNING", f"{source.name}：指定的 Sheet 不存在：{name}")
        records: list[Record] = []
        for sheet_name in selected:
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows(values_only=True)
            for _ in range(max(0, source.header_row - 1)):
                next(rows, None)
            header_values = next(rows, None)
            if not header_values:
                self.logger("WARNING", f"{source.name}/{sheet_name}：空 Sheet，已跳过")
                continue
            headers, selectors = self._headers(list(header_values))
            count = 0
            for row_number, values in enumerate(rows, start=source.header_row + 1):
                if not values or not any(value not in (None, "") for value in values):
                    continue
                raw = self._cells(list(values), headers, selectors)
                mapped = canonicalize(raw, self.aliases)
                digest = self._hash(sid, sheet_name, row_number, mapped)
                records.append(Record(source.id, source.name, sid, sheet_name, row_number, mapped, digest))
                count += 1
            self.logger("INFO", f"{source.name}/{sheet_name}：读取 {count} 行")
        workbook.close()
        return records

    def _read_private(self, source: SourceConfig, sid: str, credential: str | list[str]) -> list[Record]:
        pool = CredentialPool(credential if isinstance(credential, list) else [credential], self.logger)
        metadata = pool.call(
            lambda client: client.http_client.fetch_sheet_metadata(
                sid, params={"fields": "sheets.properties(title,hidden)"}
            )
        )
        all_names = [item["properties"]["title"] for item in metadata.get("sheets", [])]
        excludes = [*self.global_excludes, *source.exclude_sheets]
        selected, missing = choose_sheets(all_names, source.include_sheets, excludes)
        for name in missing:
            self.logger("WARNING", f"{source.name}：指定的 Sheet 不存在：{name}")
        records: list[Record] = []
        values_by_sheet: dict[str, list[list[str]]] = {}
        for start in range(0, len(selected), 100):
            names = selected[start:start + 100]
            ranges = [f"'{name.replace(chr(39), chr(39) * 2)}'!A:ZZZ" for name in names]
            response = pool.call(
                lambda client, ranges=ranges: client.http_client.values_batch_get(
                    sid,
                    ranges,
                    params={"majorDimension": "ROWS", "valueRenderOption": "FORMATTED_VALUE"},
                )
            )
            returned = response.get("valueRanges", [])
            for index, sheet_name in enumerate(names):
                values_by_sheet[sheet_name] = returned[index].get("values", []) if index < len(returned) else []
            self.logger("INFO", f"{source.name}：已批量读取 {len(names)} 个子 Sheet（1 次 API 请求）")
        for sheet_name in selected:
            values = values_by_sheet.get(sheet_name, [])
            if len(values) < source.header_row:
                self.logger("WARNING", f"{source.name}/{sheet_name}：空 Sheet，已跳过")
                continue
            headers, selectors = self._headers(values[source.header_row - 1])
            count = 0
            for row_number, row_values in enumerate(values[source.header_row:], start=source.header_row + 1):
                if not any(str(value).strip() for value in row_values):
                    continue
                raw = self._cells(list(row_values), headers, selectors)
                mapped = canonicalize(raw, self.aliases)
                digest = self._hash(sid, sheet_name, row_number, mapped)
                records.append(Record(source.id, source.name, sid, sheet_name, row_number, mapped, digest))
                count += 1
            self.logger("INFO", f"{source.name}/{sheet_name}：读取 {count} 行")
        return records

    @staticmethod
    def _gspread_client(credential: str):
        import gspread

        path = Path(credential).expanduser()
        if not path.is_file() or path.suffix.lower() != ".json":
            raise FileNotFoundError(f"服务账号 JSON 无效或不存在：{credential}")
        return gspread.service_account(filename=str(path))

    @staticmethod
    def _hash(sid: str, sheet: str, row_number: int, values: dict[str, str]) -> str:
        raw = json.dumps([sid, sheet, row_number, values], ensure_ascii=False, sort_keys=True)
        return hashlib.sha256(raw.encode("utf-8")).hexdigest()
